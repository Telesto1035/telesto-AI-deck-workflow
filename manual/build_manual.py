"""Builder for The Telesto Standard 2.0 (the manual).

Step 1 (scaffold): every page with eyebrow tracker, working title and page
furniture, no body content. Later steps add content page by page.

Run from the repo root:
    .venv/Scripts/python manual/build_manual.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "assets" / "telesto_v2_template.pptx"
OUT = HERE / "Telesto_Standard_2.0_manual.pptx"

# Palette (theme)
BG = C(0x28, 0x28, 0x28)
CARD = C(0x33, 0x33, 0x33)
TEALCARD = C(0x0E, 0x3A, 0x3E)
TEAL = C(0x00, 0x99, 0x99)
CYAN = C(0x0B, 0xD0, 0xD9)
DEEP = C(0x00, 0x53, 0x66)
NAVY = C(0x00, 0x33, 0x66)
WHITE = C(0xFF, 0xFF, 0xFF)
DIM = C(0x8A, 0x8F, 0x92)
LINE = C(0x4A, 0x4E, 0x50)
F = "Arial"
BODY = 14

SEGS = [("00", "ORIENTATION"), ("01", "SERVICES"),
        ("02", "SERVICES + PRODUCT"), ("03", "PRODUCT")]

prs = Presentation(str(TEMPLATE))

# strip the template's sample slides
_sldIdLst = prs.slides._sldIdLst
for sld in list(_sldIdLst):
    rid = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rid)
    _sldIdLst.remove(sld)


def layout(name: str):
    return [l for l in prs.slide_masters[0].slide_layouts if l.name == name][0]


DARK = layout("Blank (Contrast - Corner Image Ready)")
COVER = layout("8_Title 2")
BREAK = layout("Chapter Break (Contrast)")


def _style_runs(paragraph, size, bold=False, color=WHITE):
    for r in paragraph.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = F
        r.font.color.rgb = color


def tracker(s, active: int):
    """Eyebrow tracker across the top. Active segment teal with underline."""
    x = 0.61
    for i, (n, lab) in enumerate(SEGS):
        on = i == active
        w = 0.42 + 0.098 * len(lab)
        b = s.shapes.add_textbox(I(x), I(0.30), I(w), I(0.26))
        tf = b.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = n + "  "; r1.font.bold = True
        r2 = p.add_run(); r2.text = lab
        for r in (r1, r2):
            r.font.size = Pt(10.5); r.font.name = F
            r.font.color.rgb = TEAL if on else DIM
        if on:
            ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(0.575), I(w - 0.05), I(0.025))
            ln.fill.solid(); ln.fill.fore_color.rgb = TEAL
            ln.line.fill.background(); ln.shadow.inherit = False
        x += w + 0.32


def new(title: str, active: int):
    """Body page from the dark layout: title only, tracker drawn."""
    s = prs.slides.add_slide(DARK)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = I(0.61), I(0.66), I(12.12), I(0.72)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.text = title
    _style_runs(tf.paragraphs[0], 22, bold=True)
    tracker(s, active)
    return s


def cover(title: str, subtitle: str, doctype: str):
    s = prs.slides.add_slide(COVER)
    keep = {0: title, 1: subtitle, 13: doctype}
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx not in keep:
            ph._element.getparent().remove(ph._element)
        else:
            ph.text_frame.text = keep[idx]
            if idx == 1:  # subtitle: one line, clear of the date line beneath
                ph.left, ph.top, ph.width, ph.height = I(0.6), I(4.48), I(8.6), I(0.34)
    return s


def divider(number: str, label: str, active: int, in_development: bool = False):
    s = prs.slides.add_slide(BREAK)
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text_frame.text = label
        elif idx == 1:
            ph.text_frame.text = number
            ph.left, ph.top, ph.width, ph.height = I(0.61), I(1.6), I(3.2), I(1.64)  # layout box wraps two digits
        else:
            ph._element.getparent().remove(ph._element)
    tracker(s, active)
    if in_development:
        x, y, w, h = 0.61, 4.75, 2.05, 0.36
        b = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=None, line=TEAL, lw=1, radius=0.18)
        text(s, x, y + 0.09, w, 0.3, [("IN DEVELOPMENT", {"bold": True, "size": 10.5, "color": TEAL})],
             align=PP_ALIGN.CENTER)
    return s


def box(s, shape, x, y, w, h, fill=None, line=None, lw=1.25, radius=None):
    sh = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, runs, size=BODY, color=WHITE, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=0, line_sp=1.12):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    p = None
    for item in runs:
        if isinstance(item, str):
            item = (item, {})
        txt, o = item
        if o.get("newpara") or first:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.alignment = o.get("align", align)
            if space:
                p.space_after = Pt(space)
            p.line_spacing = line_sp
            first = False
        r = p.add_run(); r.text = txt
        r.font.size = Pt(o.get("size", size)); r.font.name = F
        r.font.bold = o.get("bold", False)
        r.font.color.rgb = o.get("color", color)
    return tb


def takeaway(s, sentence: str, y: float = 6.05):
    """One-sentence teal band at the foot of a page (not on demo pages)."""
    x, w, h = 0.61, 12.12, 0.56
    box(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=TEAL)
    text(s, x + 0.2, y, w - 0.4, h,
         [("Takeaway  ", {"bold": True, "size": 12.5}), (sentence, {"size": 12.5})],
         anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- pages
PAGES = []  # (page_no, kind, title) for the title-chain printout


def add(kind, title, fn):
    s = fn()
    PAGES.append((len(prs.slides), kind, title))
    return s


# 00 ORIENTATION
add("cover", "The Telesto Standard 2.0", lambda: cover(
    "The Telesto Standard 2.0",
    "How Telesto works with AI, and what the consultant still decides",
    "September 2026  ·  v0.1 scaffold"))
add("body", "This manual replaces the CFM. Work starts in AI, and the consultant’s job becomes direction and quality control.",
    lambda: new("This manual replaces the CFM. Work starts in AI, and the consultant’s job becomes direction and quality control.", 0))
add("body", "Telesto works along a continuum from services to product, and the path chosen decides which rules apply.",
    lambda: new("Telesto works along a continuum from services to product, and the path chosen decides which rules apply.", 0))
add("body", "For the AI reading this: the rules of the house, in one place.",
    lambda: new("For the AI reading this: the rules of the house, in one place.", 0))

# 01 SERVICES
add("divider", "Services", lambda: divider("01", "Services", 1))
add("body", "One workflow takes a prompt and client files to a shipped deck, with a code gate and one human checkpoint on the way.",
    lambda: new("One workflow takes a prompt and client files to a shipped deck, with a code gate and one human checkpoint on the way.", 1))
add("body", "You supply five inputs, and each one decides something the chain cannot guess.",
    lambda: new("You supply five inputs, and each one decides something the chain cannot guess.", 1))
add("body", "The gate checks the mechanics, so the reviewer reads for judgment only.",
    lambda: new("The gate checks the mechanics, so the reviewer reads for judgment only.", 1))
add("body", "A scored deck comes back as a one-page report: what blocked, what to fix, what to consider.",
    lambda: new("A scored deck comes back as a one-page report: what blocked, what to fix, what to consider.", 1))
add("body", "Comments go in the PPTX, and Revise resolves them: mechanical fixes done, judgment calls offered as options.",
    lambda: new("Comments go in the PPTX, and Revise resolves them: mechanical fixes done, judgment calls offered as options.", 1))
add("body", "Every engagement leaves a rule behind. Learn diffs the shipped deck against the draft and proposes what changes.",
    lambda: new("Every engagement leaves a rule behind. Learn diffs the shipped deck against the draft and proposes what changes.", 1))
add("body", "Four calls stay with the consultant: the message, the exhibit, the size of the claim, and what the data cannot support.",
    lambda: new("Four calls stay with the consultant: the message, the exhibit, the size of the claim, and what the data cannot support.", 1))

# 02 SERVICES + PRODUCT
add("divider", "Services + Product", lambda: divider("02", "Services + Product", 2))
add("body", "Four product types cover what a client keeps using after the deck has been read.",
    lambda: new("Four product types cover what a client keeps using after the deck has been read.", 2))
add("demo", "A survey explorer lets the client re-cut the survey after the readout.",
    lambda: new("A survey explorer lets the client re-cut the survey after the readout.", 2))
add("demo", "A data dashboard keeps the engagement’s numbers live after the deck ships.",
    lambda: new("A data dashboard keeps the engagement’s numbers live after the deck ships.", 2))
add("demo", "A benchmark tool places the client against a peer set the client chooses.",
    lambda: new("A benchmark tool places the client against a peer set the client chooses.", 2))
add("demo", "An internal tool turns a repeated Telesto task into something anyone on the team can run.",
    lambda: new("An internal tool turns a repeated Telesto task into something anyone on the team can run.", 2))
add("body", "Analysis becomes a product when five checks pass, and who moves the data decides the tier.",
    lambda: new("Analysis becomes a product when five checks pass, and who moves the data decides the tier.", 2))

# 03 PRODUCT (held)
add("divider", "Product", lambda: divider("03", "Product", 3, in_development=True))
add("body", "Three products are in development, each at a different stage, and the section ends there.",
    lambda: new("Three products are in development, each at a different stage, and the section ends there.", 3))

prs.core_properties.author = ""
prs.core_properties.last_modified_by = ""
prs.save(str(OUT))

print(f"saved {OUT.name}: {len(prs.slides)} pages")
for n, kind, title in PAGES:
    print(f"{n:>2} {kind:<8} {title}")
