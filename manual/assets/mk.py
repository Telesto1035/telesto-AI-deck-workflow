from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG=C(0x28,0x28,0x28); CARD=C(0x33,0x33,0x33); CARD2=C(0x3B,0x3B,0x3B)
TEALCARD=C(0x0E,0x3A,0x3E)
TEAL=C(0x00,0x99,0x99); CYAN=C(0x0B,0xD0,0xD9); WHITE=C(0xFF,0xFF,0xFF)
DIM=C(0x8A,0x8F,0x92); LINE=C(0x4A,0x4E,0x50)
F="Arial"; BODY=14

prs=Presentation('/home/claude/vs2.pptx')
# strip existing slides
xml_slides=prs.slides._sldIdLst
for sld in list(xml_slides):
    rid=sld.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rid); xml_slides.remove(sld)
DARK=[l for l in prs.slide_masters[0].slide_layouts if l.name=='Blank (Contrast - Corner Image Ready)'][0]

def new(title, active):
    s=prs.slides.add_slide(DARK)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx!=0: ph._element.getparent().remove(ph._element)
    t=s.shapes.title
    t.left,t.top,t.width,t.height=I(0.61),I(0.66),I(12.12),I(0.72)
    tf=t.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP; tf.text=title
    for r in tf.paragraphs[0].runs:
        r.font.size=Pt(22); r.font.bold=True; r.font.name=F; r.font.color.rgb=WHITE
    # eyebrow tracker
    segs=[("00","ORIENTATION"),("01","SERVICES"),("02","SERVICES + PRODUCT"),("03","PRODUCT")]
    x=0.61
    for i,(n,lab) in enumerate(segs):
        on = (i==active)
        w = 0.42+0.098*len(lab)
        b=s.shapes.add_textbox(I(x),I(0.30),I(w),I(0.26)); tf2=b.text_frame
        tf2.word_wrap=False; tf2.margin_left=tf2.margin_right=tf2.margin_top=tf2.margin_bottom=0
        p=tf2.paragraphs[0]
        r1=p.add_run(); r1.text=n+"  "; r1.font.bold=True
        r2=p.add_run(); r2.text=lab
        for r in (r1,r2):
            r.font.size=Pt(10.5); r.font.name=F; r.font.color.rgb= TEAL if on else DIM
        if on:
            ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,I(x),I(0.575),I(w-0.05),I(0.025))
            ln.fill.solid(); ln.fill.fore_color.rgb=TEAL; ln.line.fill.background(); ln.shadow.inherit=False
        x+=w+0.32
    return s

def box(s, shape, x,y,w,h, fill=None, line=None, lw=1.25, radius=None):
    sh=s.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0]=radius
    sh.text_frame.text=""
    return sh

def text(s,x,y,w,h,runs,size=BODY,color=WHITE,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=0,line_sp=1.12):
    tb=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=anchor
    first=True
    for item in runs:
        if isinstance(item,str): item=(item,{})
        txt,o=item
        if o.get('newpara') or first:
            p=tf.paragraphs[0] if first else tf.add_paragraph()
            p.alignment=o.get('align',align)
            if space: p.space_after=Pt(space)
            p.line_spacing=line_sp
            first=False
        r=p.add_run(); r.text=txt
        r.font.size=Pt(o.get('size',size)); r.font.name=F
        r.font.bold=o.get('bold',False)
        r.font.color.rgb=o.get('color',color)
    return tb

# ============ P3 : CONTINUUM ============
s=new("Telesto works along a continuum from services to product, and the path chosen decides which rules apply.",0)
CH=[("01","SERVICES","Project-based consulting"),
    ("02","SERVICES + PRODUCT","Consulting with a data product"),
    ("03","PRODUCT","The product is the deliverable")]
cw,step,cy,chh=4.10,3.98,1.85,0.82
for i,(n,lab,sub) in enumerate(CH):
    x=0.61+step*i
    shape=MSO_SHAPE.PENTAGON if i==0 else MSO_SHAPE.CHEVRON
    fill=[TEAL,C(0x00,0x73,0x73),C(0x00,0x50,0x52)][i]
    sh=box(s,shape,x,cy,cw,chh,fill=fill)
    sh.adjustments[0]=0.22
    text(s,x+(0.30 if i==0 else 0.62),cy+0.19,cw-0.9,0.5,
         [(n+"   ",{'bold':True,'size':15,'color':WHITE}),(lab,{'bold':True,'size':15,'color':WHITE})])
# cards
cardw,cardx=3.89,[0.61,4.72,8.83]
cyy,cardh=2.95,3.20
BODYTXT=[
 [("Analysis, findings and recommendations are delivered as a deck.",{})],
 [("A deck states the findings; a product the client keeps using answers the questions that follow.",{})],
 [("A live application the client operates. A short deck exists to explain it.",{})],
]
EX=[["Intelligence programs","Strategy and PMO engagements"],
    ["Client chatbot","Public sector market tools"],
    ["Standalone application","Sold on its own terms"]]
for i in range(3):
    x=cardx[i]
    box(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,cyy,cardw,cardh,fill=CARD,radius=0.04)
    text(s,x+0.26,cyy+0.30,cardw-0.52,1.0,BODYTXT[i])
    ln=box(s,MSO_SHAPE.RECTANGLE,x+0.26,cyy+1.42,cardw-0.52,0.014,fill=LINE)
    text(s,x+0.26,cyy+1.62,cardw-0.52,0.3,[("IN PRACTICE",{'bold':True,'size':11,'color':TEAL})])
    yy=cyy+1.96
    for e in EX[i]:
        d=box(s,MSO_SHAPE.OVAL,x+0.26,yy+0.075,0.11,0.11,fill=TEAL)
        text(s,x+0.52,yy,cardw-0.80,0.34,[(e,{})])
        yy+=0.42
    if i==2:
        tag=box(s,MSO_SHAPE.ROUNDED_RECTANGLE,x+cardw-1.72,cyy+0.0-0.0,1.5,0.0,fill=None)
        tag._element.getparent().remove(tag._element)
        box(s,MSO_SHAPE.ROUNDED_RECTANGLE,x+0.26,cyy+cardh-0.44,1.92,0.36,fill=None,line=TEAL,lw=1,radius=0.18)
        text(s,x+0.26,cyy+cardh-0.345,1.92,0.3,[("IN DEVELOPMENT",{'bold':True,'size':10.5,'color':TEAL})],align=PP_ALIGN.CENTER)

# ============ P4 : WORKFLOW ============
s=new("Every deliverable starts with a one-slide brief, is drafted with AI, and clears a self-check before partner review.",0)
NODES=[("INTAKE","Pick the path: services, services and product, or product"),
       ("BRIEF","One slide: goal, audience, value at stake, storyline"),
       ("BUILD","AI drafts from the approved brief"),
       ("QC BOT","The consultant scores the deck and fixes what it flags"),
       ("REVIEW","Partner reads for judgment, then it ships")]
nw,gap,ny,nh=2.16,0.30,1.72,1.62
for i,(lab,desc) in enumerate(NODES):
    x=0.61+(nw+gap)*i
    gate = (i==1); check = (i==3)
    box(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,ny,nw,nh,
        fill=(TEALCARD if gate else CARD),line=(TEAL if (gate or check) else None),lw=1.25,radius=0.05)
    text(s,x+0.20,ny+0.22,nw-0.40,0.3,[(lab,{'bold':True,'size':12,'color':(TEAL if (gate or check) else WHITE)})])
    text(s,x+0.20,ny+0.62,nw-0.40,nh-0.8,[(desc,{'size':12.5})])
    if i<4:
        box(s,MSO_SHAPE.RIGHT_ARROW,x+nw+0.06,ny+nh/2-0.09,0.18,0.18,fill=DIM)
# single gate under BRIEF
gx=0.61+(nw+gap)*1
box(s,MSO_SHAPE.ROUNDED_RECTANGLE,gx,ny+nh+0.16,nw,0.40,fill=None,line=TEAL,lw=1,radius=0.18)
text(s,gx,ny+nh+0.26,nw,0.3,[("GATE",{'bold':True,'size':10.5,'color':TEAL}),("   Green light to build",{'size':10.5,'color':WHITE})],align=PP_ALIGN.CENTER)
# loop back: QC BOT -> BUILD
bx=0.61+(nw+gap)*2; qx=0.61+(nw+gap)*3
xb=bx+nw/2; xq=qx+nw/2; yl=ny+nh+0.30
box(s,MSO_SHAPE.RECTANGLE,xq-0.008,ny+nh,0.016,yl-(ny+nh),fill=TEAL)
box(s,MSO_SHAPE.RECTANGLE,xb-0.008,ny+nh+0.20,0.016,yl-(ny+nh+0.20),fill=TEAL)
box(s,MSO_SHAPE.RECTANGLE,xb,yl,xq-xb,0.016,fill=TEAL)
box(s,MSO_SHAPE.UP_ARROW,xb-0.09,ny+nh+0.04,0.18,0.18,fill=TEAL)
text(s,xb,yl+0.10,xq-xb,0.28,[("Flags come back, fix and re-run",{'size':11.5,'color':WHITE})],align=PP_ALIGN.CENTER)
# scored dimensions
hy=4.30
text(s,0.61,hy,8.0,0.3,[("WHAT THE QC BOT SCORES, 1 TO 5 PER DECK",{'bold':True,'size':11,'color':TEAL})])
DIMS=[("outcome","Clarity of outcome and recommendation","One takeaway, and the reader knows what to decide"),
      ("flow","Logical flow","Parts align, the order holds, no missing pages"),
      ("value","Financial value at stake","Dollar impact stated and defensible"),
      ("accuracy","Technical accuracy","Every number traced to a real source")]
ty=4.68; th=1.70; tw=(12.10-0.60)/4
for i,(ic,name,desc) in enumerate(DIMS):
    x=0.61+(tw+0.20)*i
    box(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,ty,tw,th,fill=CARD,radius=0.05)
    s.shapes.add_picture('/home/claude/icons/%s.png'%ic,I(x+0.24),I(ty+0.22),I(0.34),I(0.34))
    text(s,x+0.24,ty+0.62,tw-0.48,0.52,[(name,{'bold':True,'size':12.5})],anchor=MSO_ANCHOR.BOTTOM)
    text(s,x+0.24,ty+1.22,tw-0.48,0.44,[(desc,{'size':12})])

# ============ P5 : SKILLS AND TOOLS ============
s=new("Each path uses a different mix of skills and tools.",0)
LX,LW=0.61,1.75
CXS=[2.50,5.95,9.40]; CW=3.31
hy=1.58; hh=0.50
HEAD=[("01","SERVICES"),("02","SERVICES + PRODUCT"),("03","PRODUCT")]
for i,(n,lab) in enumerate(HEAD):
    fill=[TEAL,C(0x00,0x73,0x73),C(0x00,0x50,0x52)][i]
    box(s,MSO_SHAPE.ROUNDED_RECTANGLE,CXS[i],hy,CW,hh,fill=fill,radius=0.10)
    text(s,CXS[i]+0.22,hy+0.13,CW-0.44,0.3,[(n+"   ",{'bold':True,'size':12.5}),(lab,{'bold':True,'size':12.5})])
    if i==2:
        text(s,CXS[i]+0.22,hy+0.15,CW-0.44,0.3,[("IN DEVELOPMENT",{'bold':True,'size':10,'color':WHITE})],align=PP_ALIGN.RIGHT)
ROWS=[("Consultant skills",2.22,1.62,
       ["Storyline and synthesis","Executive writing","Client psychology"],
       ["Everything at left, plus:","Data structuring","Dashboard and tool design","Client data handling"],
       ["Everything at left, plus:","Product specification","User experience","Deployment and support"]),
      ("Build tools",3.98,1.62,
       ["Claude with Telesto skills","PowerPoint"],
       ["Everything at left, plus:","Claude Code","Replit","Google Cloud for live data"],
       ["Everything at left, plus:","GitHub","Managed hosting","Identity and access"]),
      ("Where the work lives",5.74,0.72,
       ["A deck file"],["A client portal, with the deck"],["A live application"])]
for label,ry,rh,a,b,c in ROWS:
    box(s,MSO_SHAPE.ROUNDED_RECTANGLE,LX,ry,LW,rh,fill=C(0x14,0x33,0x35),radius=0.06)
    text(s,LX+0.18,ry,LW-0.36,rh,[(label,{'bold':True,'size':12.5,'color':TEAL})],anchor=MSO_ANCHOR.MIDDLE)
    for i,items in enumerate((a,b,c)):
        box(s,MSO_SHAPE.ROUNDED_RECTANGLE,CXS[i],ry,CW,rh,fill=CARD,radius=0.06)
        runs=[]
        for j,it in enumerate(items):
            lead = it.endswith('plus:')
            runs.append((it,{'newpara':j>0,'size':12.5,'color':(TEAL if lead else WHITE)}))
        text(s,CXS[i]+0.22,ry+0.20,CW-0.44,rh-0.36,runs,space=4)


prs.core_properties.author=""
prs.core_properties.last_modified_by=""
prs.save('/home/claude/Telesto_Standard_v2_P3-P5.pptx')
print("saved")
