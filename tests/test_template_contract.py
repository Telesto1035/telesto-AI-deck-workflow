"""Template contract conformance (qc_check.py --template; draft-deck references/template-contract.md).

Built templates check the contract logic everywhere. Client templates never enter the repo: the registry test
runs against the templates listed in tests/fixtures/client/templates.json (gitignored) and skips without it."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
SCRIPTS = REPO / ".claude" / "skills" / "deck-qc" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import qc_check  # noqa: E402

REGISTRY = HERE / "fixtures" / "client" / "templates.json"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


@pytest.fixture(scope="module")
def rules():
    return qc_check.load_rules(None)


def _template(tmp_path, pages, sections=None, name="Client_Template_2026_v1.pptx"):
    """pages: list of (shape name, text). sections: list of (section name, [1-based slide numbers])."""
    prs = Presentation()
    for shape_name, text in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(0.5))
        tb.name = shape_name
        tb.text_frame.text = text
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.size = Pt(12)
    if sections:
        ids = [s.slide_id for s in prs.slides]
        root = prs.part._element
        ext_lst = root.find("{%s}extLst" % P)
        if ext_lst is None:
            ext_lst = etree.SubElement(root, "{%s}extLst" % P)
        ext = etree.SubElement(ext_lst, "{%s}ext" % P, uri=qc_check.SECTION_URI)
        lst = etree.SubElement(ext, "{%s}sectionLst" % P14, nsmap={"p14": P14})
        for i, (sec_name, slides) in enumerate(sections):
            sec = etree.SubElement(lst, "{%s}section" % P14, name=sec_name,
                                   id="{00000000-0000-0000-0000-%012d}" % i)
            sl = etree.SubElement(sec, "{%s}sldIdLst" % P14)
            for sno in slides:
                etree.SubElement(sl, "{%s}sldId" % P14, id=str(ids[sno - 1]))
    out = tmp_path / name
    prs.save(str(out))
    return out


STARTER = ("Title 1", "Starter page title")
EXEMPLAR = ("Stamp", "EXEMPLAR · DUMMY DATA")
README = ("Eyebrow", "TEMPLATE README · NOT A DECK PAGE · Template contract v1")


def _blocks(findings):
    return [f for f in findings if f.severity == "block"]


def test_conforming_template_passes(tmp_path, rules):
    path = _template(tmp_path, [STARTER, EXEMPLAR, README],
                     [("Starter", [1]), ("Exemplar", [2]), ("Readme", [3])])
    _d, findings = qc_check.check_template_contract(path, rules)
    assert not [f for f in findings if f.severity in ("block", "warn")], [f.detail for f in findings]


def test_readme_without_contract_version_blocks(tmp_path, rules):
    path = _template(tmp_path, [STARTER, EXEMPLAR, ("Eyebrow", "TEMPLATE README · NOT A DECK PAGE")])
    _d, findings = qc_check.check_template_contract(path, rules)
    assert any("contract version" in f.detail for f in _blocks(findings))


def test_missing_readme_blocks(tmp_path, rules):
    _d, findings = qc_check.check_template_contract(_template(tmp_path, [STARTER, EXEMPLAR]), rules)
    assert any("no readme page" in f.detail for f in _blocks(findings))


def test_older_contract_version_warns(tmp_path, rules):
    r = dict(rules, template_contract_version=2)
    _d, findings = qc_check.check_template_contract(_template(tmp_path, [STARTER, EXEMPLAR, README]), r)
    assert not _blocks(findings)
    assert any("meets template contract v1" in f.detail for f in findings if f.severity == "warn")


def test_exemplar_section_page_without_stamp_blocks(tmp_path, rules):
    """The case sections exist for: an exemplar page whose stamp was deleted reads as a starter page."""
    path = _template(tmp_path, [STARTER, ("Title 1", "Real client content"), README],
                     [("Starter", [1]), ("Exemplar", [2]), ("Readme", [3])])
    _d, findings = qc_check.check_template_contract(path, rules)
    assert [f.slide for f in _blocks(findings)] == [2]


def test_starter_source_line_outside_a_named_home_warns(tmp_path, rules):
    path = _template(tmp_path, [("TextBox 7", "Source: [system], accessed [month year]"), EXEMPLAR, README])
    _d, findings = qc_check.check_template_contract(path, rules)
    assert any(f.slide == 1 and "not named Source" in f.detail for f in findings if f.severity == "warn")


def test_cli_template_mode_exit_code(tmp_path):
    good = _template(tmp_path, [STARTER, EXEMPLAR, README], name="good_template.pptx")
    bad = _template(tmp_path, [STARTER, EXEMPLAR], name="bad_template.pptx")
    assert qc_check.main([str(good), "--template"]) == 0
    assert qc_check.main([str(bad), "--template"]) == 1


def _registry():
    if not REGISTRY.exists():
        return []
    entries = json.loads(REGISTRY.read_text(encoding="utf-8")).get("templates", [])
    return [pytest.param(e, id=e.get("client", Path(e["path"]).stem)) for e in entries]


@pytest.mark.skipif(not REGISTRY.exists(), reason="no local template registry (tests/fixtures/client/templates.json)")
@pytest.mark.parametrize("entry", _registry())
def test_registered_templates_meet_the_contract(entry, rules):
    path = Path(entry["path"]).expanduser()
    if not path.exists():
        pytest.skip("template not on this machine: %s" % path)
    _d, findings = qc_check.check_template_contract(path, rules)
    assert not _blocks(findings), [(f.slide, f.detail) for f in _blocks(findings)]
