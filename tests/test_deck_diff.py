"""Tests for scripts/deck_diff.py: matching, classification, appendix handling and the comment join."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(HERE))

import deck_diff  # noqa: E402
import plant_comments  # noqa: E402

CLEAN = HERE / "fixtures" / "Telesto_Standard_v2_P3-P5.pptx"


def _edit(src: Path, dst: Path, title_suffix: str = "", number_from: str | None = None, number_to: str = "") -> Path:
    """Copy src to dst with the first slide's title changed and, optionally, one number rewritten on slide 2."""
    prs = Presentation(str(src))
    t = prs.slides[0].shapes.title
    if t is not None and title_suffix:
        t.text_frame.paragraphs[0].runs[0].text += title_suffix
    if number_from and len(prs.slides) > 1:
        for sh in prs.slides[1].shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if number_from in r.text:
                            r.text = r.text.replace(number_from, number_to)
    prs.save(str(dst))
    return dst


@pytest.fixture(scope="module")
def pair(tmp_path_factory):
    d = tmp_path_factory.mktemp("diff")
    before = plant_comments.plant(CLEAN, d / "deck_v1.pptx")
    after = _edit(before, d / "deck_v2.pptx", title_suffix=" and one more clause for the reviewer")
    return before, after


def test_identical_decks_report_no_changes(tmp_path, capsys):
    assert deck_diff.main([str(CLEAN), str(CLEAN)]) == 0
    out = capsys.readouterr().out
    assert "| number | 0 |" in out and "| structure | 0 |" in out and "## Changes by slide" in out


def test_title_edit_is_a_style_change_and_pages_stay_matched(pair, tmp_path):
    before, after = pair
    js = tmp_path / "d.json"
    assert deck_diff.main([str(before), str(after), "--json", str(js)]) == 0
    res = json.loads(js.read_text(encoding="utf-8"))
    assert res["removed"] == [] and res["added"] == []
    first = next(p for p in res["pairs"] if p["after_slide"] == 1)
    kinds = {(c["kind"], c["class"]) for c in first["changes"]}
    assert ("title", "style") in kinds


def test_number_regex_ignores_sentence_full_stop():
    assert deck_diff.numbers("grew to 2024.") == ["2024"]
    assert deck_diff.classify("Revenue 2024.", "Revenue 2024", is_title=False) == "style"
    assert deck_diff.classify("Revenue 2024", "Revenue 2025", is_title=False) == "number"


def test_comments_join_and_ledger_rows(pair, capsys):
    before, after = pair
    assert deck_diff.main([str(before), str(after), "--comments", "--ledger-rows", "--date", "2026-09-13"]) == 0
    out = capsys.readouterr().out
    assert "## Comments in BEFORE and what changed on their slides" in out
    assert "### 1. slide 1 -> 1 (Alex Kruzel, active)" in out
    assert "> Make this title more descriptive" in out
    assert "> reply, Alex Kruzel: Agreed, will do" in out
    assert "no change on this slide" in out  # the legacy comment on slide 2, which was not edited
    assert "## Ledger rows" in out
    assert "| 2026-09-13 | deck | v1 | v2 | 1 |  |  | TODO: Make this title more descriptive |" in out
    assert "| 2026-09-13 | deck | v1 | v2 | 3 |  |  | NO CHANGE ON THIS SLIDE: Please add the footnotes |" in out


def test_appendix_pages_are_marked():
    slides = deck_diff.load(CLEAN)
    assert all("appendix" in s for s in slides)
    assert deck_diff.APPENDIX_RE.match("Appendix: crosstabs") and deck_diff.APPENDIX_RE.match("Crosstabs - AI")
    assert not deck_diff.APPENDIX_RE.match("Backups are not the point")
